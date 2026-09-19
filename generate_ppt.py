from pptx import Presentation

def create_presentation():
    prs = Presentation()
    
    # Define slide layouts
    title_layout = prs.slide_layouts[0]
    bullet_layout = prs.slide_layouts[1]
    
    # Title Slide
    slide = prs.slides.add_slide(title_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "FinFlow"
    subtitle.text = "Personal Finance Intelligence for Young Earners"
    
    # Slide 2
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "The Problem"
    tf = slide.placeholders[1].text_frame
    tf.text = "Financial ambiguity for young earners:"
    tf.add_paragraph().text = "• Irregular income from stipends, freelancing, or entry-level jobs."
    tf.add_paragraph().text = "• Difficulty tracking where money goes across many micro-transactions."
    tf.add_paragraph().text = "• Uncertainty about upcoming monthly cash flow."
    
    # Slide 3
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "The Solution: Three Key Questions"
    tf = slide.placeholders[1].text_frame
    tf.text = "FinFlow is built to answer:"
    tf.add_paragraph().text = "1. Where did my money go? (Dashboard & categorical breakdowns)"
    tf.add_paragraph().text = "2. Is anything wrong? (ML-powered fraud alerts & budget limits)"
    tf.add_paragraph().text = "3. Will I be okay next month? (Cash flow forecasting & subscription audit)"
    
    # Slide 4
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "Key Features Overview"
    tf = slide.placeholders[1].text_frame
    tf.text = "Core functionalities:"
    tf.add_paragraph().text = "• Interactive Dashboard: Real-time financial health score."
    tf.add_paragraph().text = "• Smart Budgets: Track spending across categories with dynamic limits."
    tf.add_paragraph().text = "• Subscriptions Tracker: Identify recurring costs and hidden fees."
    tf.add_paragraph().text = "• Financial Goals: Measure actual savings rates against long-term targets."
    
    # Slide 5
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "System Architecture"
    tf = slide.placeholders[1].text_frame
    tf.text = "End-to-end data pipeline:"
    tf.add_paragraph().text = "1. Data Generator: Produces realistic mock transactions."
    tf.add_paragraph().text = "2. ETL Pipeline: Cleans and categorizes raw data."
    tf.add_paragraph().text = "3. Database: PostgreSQL with a robust Star Schema."
    tf.add_paragraph().text = "4. Machine Learning: scikit-learn models for inference."
    tf.add_paragraph().text = "5. Application Layer: Flask Web App serving HTML and REST API."
    
    # Slide 6
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "Data Engineering (ETL)"
    tf = slide.placeholders[1].text_frame
    tf.text = "Handling the data at scale:"
    tf.add_paragraph().text = "• Star Schema: Organized into fact and dimension tables (Accounts, Merchants, Dates)."
    tf.add_paragraph().text = "• Transformation: Detects recurring merchants, standardizes dates."
    tf.add_paragraph().text = "• Categorization: Assigns expenses using keyword/MCC mapping."
    tf.add_paragraph().text = "• Scalable: Uses pandas and psycopg2 for bulk upserts."

    # Slide 7
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "Machine Learning Integrations"
    tf = slide.placeholders[1].text_frame
    tf.text = "Intelligent backend systems:"
    tf.add_paragraph().text = "• Fraud Detection: Isolation Forest algorithm flags anomalous amounts, times, and merchants."
    tf.add_paragraph().text = "• Cash Flow Forecast: Linear Regression predicts 30-day spending trends."
    tf.add_paragraph().text = "• Evaluation: Models are backtested against naive baselines (e.g. 30-day trailing averages)."

    # Slide 8
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "Security & API"
    tf = slide.placeholders[1].text_frame
    tf.text = "Robust and protected:"
    tf.add_paragraph().text = "• API Layer: JSON REST API under `/api/v1/*`."
    tf.add_paragraph().text = "• Authentication: Secure sessions with Flask-Login, plus CSRF tokens on all forms."
    tf.add_paragraph().text = "• Isolation: All queries are strictly scoped to the authenticated user's ID."

    # Slide 9
    slide = prs.slides.add_slide(bullet_layout)
    title = slide.shapes.title
    title.text = "Current Roadmap & Enhancements"
    tf = slide.placeholders[1].text_frame
    tf.text = "Moving to Production Quality:"
    tf.add_paragraph().text = "• Model Calibration: Persisting score thresholds at training time for consistent alerts."
    tf.add_paragraph().text = "• Security Patching: Mitigating SQL injection in analytical endpoints."
    tf.add_paragraph().text = "• Better Telemetry: Retaining granular timestamps in the ETL for better ML feature engineering."
    
    # Slide 10
    slide = prs.slides.add_slide(title_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Conclusion"
    subtitle.text = "Thank you! Any Questions?"

    # Save presentation
    prs.save("FinFlow_Presentation.pptx")

if __name__ == "__main__":
    create_presentation()
    print("Presentation created successfully as FinFlow_Presentation.pptx")
